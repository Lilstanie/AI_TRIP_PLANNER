"""Build ELEC5620_Stage1_Slides.pptx (16:9, editable, speaker script in the notes).

Usage: python3 build_deck.py [output.pptx]
Needs python-pptx. Diagrams come from assets/diagrams (re-rendered from the report's
mermaid sources in mmd-src/ with the deck theme), icons from assets/icons (Lucide, ISC).
Fonts are Georgia (titles) and Arial (body) so the deck looks the same on Windows and Mac.
"""
import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE as DASH
from pptx.oxml.ns import qn
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
DIA = os.path.join(HERE, 'assets', 'diagrams')
ICO = os.path.join(HERE, 'assets', 'icons')

NAVY, INK, MUTED, ORANGE = '14213D', '2A3550', '5A6478', 'C2571A'
PAPER, CARD, LINE, TINT, PEACH = 'F7F6F1', 'FFFFFF', 'E2E5EB', 'EEF2F8', 'FBEDE3'
HEAD, BODY = 'Georgia', 'Arial'
MEMBERS = {
    'A': ('Ziqi He', 'Orchestrator', '1F7A8C'),
    'B': ('Tingsong Jin', 'Itinerary and transport', '6B4FA0'),
    'C': ('Yi Qiao', 'Accommodation and budget', 'C2571A'),
    'D': ('Jiahang Bian', 'Destination and dining', '2E7D4F'),
    'E': ('Weihao Wang', 'Workspace and memory', '2F5DA8'),
}
TOTAL = 24
ML, MR = 0.6, 12.733          # content left / right edge (inches)
CT, CB = 1.62, 6.82           # content top / bottom

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def rgb(h):
    return RGBColor.from_string(h)


def shadow(shape, blur=9, dist=2, alpha=14):
    spPr = shape._element.spPr
    eff = etree.SubElement(spPr, qn('a:effectLst'))
    sh = etree.SubElement(eff, qn('a:outerShdw'), blurRad=str(Pt(blur)), dist=str(Pt(dist)),
                          dir='5400000', algn='t', rotWithShape='0')
    clr = etree.SubElement(sh, qn('a:srgbClr'), val='14213D')
    etree.SubElement(clr, qn('a:alpha'), val=str(alpha * 1000))


def no_effects(shape):
    etree.SubElement(shape._element.spPr, qn('a:effectLst'))


def box(s, x, y, w, h, fill=CARD, line=LINE, r=0.12, lw=0.75, shad=False, dash=None, shape=None):
    kind = shape or (MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE)
    sp = s.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = min(0.5, r / min(w, h))
    if fill:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    else:
        sp.fill.background()
    if line:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(lw)
        if dash:
            sp.line.dash_style = dash
    else:
        sp.line.fill.background()
    if shad:
        shadow(sp)
    else:
        no_effects(sp)
    sp.text_frame.text = ''
    return sp


def text(s, x, y, w, h, paras, size=14, color=INK, font=BODY, bold=False, align='l',
         anchor='t', italic=False, spacing=1.15, after=0, wrap=True):
    """paras: str | list of paragraphs; a paragraph is str or list of (text, overrides)."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {'t': MSO_ANCHOR.TOP, 'm': MSO_ANCHOR.MIDDLE, 'b': MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras, str):
        paras = [paras]
    for i, p in enumerate(paras):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = {'l': PP_ALIGN.LEFT, 'c': PP_ALIGN.CENTER, 'r': PP_ALIGN.RIGHT}[align]
        para.line_spacing = spacing
        para.space_after = Pt(after)
        runs = [(p, {})] if isinstance(p, str) else p
        for t, o in runs:
            r = para.add_run(); r.text = t
            f = r.font
            f.name = o.get('font', font); f.size = Pt(o.get('size', size))
            f.bold = o.get('bold', bold); f.italic = o.get('italic', italic)
            f.color.rgb = rgb(o.get('color', color))
    return tb


def icon(s, name, x, y, size=0.34, color=NAVY):
    path = os.path.join(ICO, f'{name}-{color}.png')
    if not os.path.exists(path):
        import cairosvg  # only needed when a new icon/colour pair is introduced
        src = open(os.path.join(os.environ.get('LUCIDE_DIR', ''), f'{name}.svg')).read()
        src = src.replace('currentColor', '#' + color).replace('stroke-width="2"', 'stroke-width="1.8"')
        cairosvg.svg2png(bytestring=src.encode(), write_to=path, output_width=192, output_height=192)
    return s.shapes.add_picture(path, Inches(x), Inches(y), Inches(size), Inches(size))


def badge(s, x, y, label, fill=ORANGE, d=0.3, size=11, fg='FFFFFF', ring=True):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
    if ring:
        c.line.color.rgb = rgb('FFFFFF'); c.line.width = Pt(1.5)
    else:
        c.line.fill.background()
    no_effects(c)
    tf = c.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(label)
    r.font.size = Pt(size); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = rgb(fg)
    return c


def pill(s, x, y, label, fill, fg='FFFFFF', size=10, h=0.3, padding=0.16, bold=True, line=None, w=None):
    w = w or (len(label) * size * 0.0078 + 2 * padding)
    b = box(s, x, y, w, h, fill=fill, line=line, r=h / 2)
    tf = b.text_frame
    tf.margin_left = tf.margin_right = Inches(0.04); tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE; tf.word_wrap = False
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = label
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = BODY; r.font.color.rgb = rgb(fg)
    return w


def figure(s, name, x, y, w, h, pad=0.14, card=True, callouts=(), anchor='c'):
    """Diagram on a white card, fitted inside (x, y, w, h). callouts: (n, fx, fy) in image fractions."""
    if card:
        box(s, x, y, w, h, shad=True)
    path = os.path.join(DIA, name)
    iw, ih = Image.open(path).size
    bw, bh = w - 2 * pad, h - 2 * pad
    sc = min(bw / iw, bh / ih)
    pw, ph = iw * sc, ih * sc
    px = x + pad + (bw - pw) / 2 if anchor == 'c' else x + pad
    py = y + pad + (bh - ph) / 2
    s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))
    for n, fx, fy in callouts:
        badge(s, px + fx * pw - 0.15, py + fy * ph - 0.15, n)
    return px, py, pw, ph


def note(s, n, x, y, w, head, body, color=ORANGE, size=12.5, h=0.75):
    badge(s, x, y + 0.01, n, fill=color, ring=False, d=0.28, size=10.5)
    paras = [[(head, {'bold': True, 'color': NAVY})]] if head else []
    if body:
        paras.append(body)
    text(s, x + 0.42, y, w - 0.42, h, paras, size=size, color=INK, after=2)


def chrome(s, idx, member, section, title):
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(PAPER)
    col = MEMBERS[member][2] if member else ORANGE
    box(s, 0, 0, 0.11, 7.5, fill=col, line=None, r=0)
    x = ML
    if member:
        name = MEMBERS[member][0]
        x += pill(s, ML, 0.38, f'{member} · {name}', col, size=10) + 0.16
    text(s, x, 0.38, 6.5, 0.3, section.upper(), size=10, color=col, bold=True, anchor='m')
    # progress: which member is speaking
    for i, m in enumerate('ABCDE'):
        on = m == member
        badge(s, MR - (5 - i) * 0.36 + 0.06, 0.38, m, fill=MEMBERS[m][2] if on else 'E7E9EE',
              fg='FFFFFF' if on else '8A93A6', d=0.3, size=10, ring=False)
    text(s, ML, 0.8, MR - ML, 0.62, title, size=26, color=NAVY, font=HEAD, bold=True, anchor='m')
    box(s, ML, 7.0, MR - ML, 0.012, fill='DDDAD0', line=None, r=0)
    text(s, ML, 7.08, 8, 0.25, 'AI Trip Planner  ·  ELEC5620 Stage 1  ·  Mon 10-12 Group 14',
         size=9, color='8A93A6')
    text(s, MR - 2, 7.08, 2, 0.25, f'{idx:02d} / {TOTAL}', size=9, color='8A93A6', align='r')


def slide(idx, member, section, title, notes):
    s = prs.slides.add_slide(BLANK)
    if section is not None:
        chrome(s, idx, member, section, title)
    s.notes_slide.notes_text_frame.text = notes
    return s


def card(s, x, y, w, h, fill=CARD, shad=True, accent=None):
    b = box(s, x, y, w, h, fill=fill, shad=shad, line=None if fill != CARD else LINE)
    if accent:
        box(s, x, y + 0.18, 0.06, h - 0.36, fill=accent, line=None, r=0.03)
    return b


def behaviour_template(idx, m):
    name, role, col = MEMBERS[m]
    s = slide(idx, m, 'Individual behaviour models', f"[{m}'s use case]: activity, sequence and state machine",
              f"[About 25 seconds. Name member {m}'s ad hoc requirement and use case, then point to where "
              "the LLM sits in each diagram and which code checks it. Replace the three placeholders with the diagrams.]")
    card(s, ML, CT, MR - ML, 0.62, accent=col)
    text(s, ML + 0.3, CT + 0.1, MR - ML - 0.5, 0.42,
         [[('Ad hoc requirement  ', {'bold': True, 'color': col}), ('[AH-' + m + '1, one line in the traveller\'s words]', {'italic': True}),
           ('      Use case  ', {'bold': True, 'color': col}), ('[UC-' + m + 'x name]', {'italic': True})]],
         size=13, anchor='m', color=MUTED)
    w = (MR - ML - 0.4) / 3
    for i, (lab, ic) in enumerate((('Activity diagram', 'workflow'), ('Sequence diagram', 'git-branch'),
                                   ('State machine', 'git-pull-request'))):
        x = ML + i * (w + 0.2)
        y = CT + 0.82
        box(s, x, y, w, CB - y, fill='FBFAF7', line='B9C0CD', lw=1.25,
            dash=DASH.DASH)
        badge(s, x + 0.22, y + 0.22, i + 1, fill=col, ring=False, d=0.3, size=11)
        text(s, x + 0.62, y + 0.2, w - 0.8, 0.34, lab, size=14, bold=True, color=NAVY, anchor='m')
        icon(s, 'image-plus', x + w / 2 - 0.3, y + 1.55, 0.6, '9AA5B8')
        text(s, x + 0.2, y + 2.3, w - 0.4, 0.9,
             ['Drop your diagram here', 'Fit to this box, keep the label row'],
             size=11.5, color='8A93A6', align='c', after=3)
        text(s, x + 0.2, CB - 0.62, w - 0.4, 0.4,
             [[('LLM step  ', {'bold': True, 'color': col}), ('[where it sits]  ', {'italic': True}),
               ('Code check  ', {'bold': True, 'color': col}), ('[what validates it]', {'italic': True})]],
             size=10, color=MUTED, align='c')


# ---------------------------------------------------------------- 1 cover
s = slide(1, None, None, '', "Hi, we are Group 14 from the Monday 10 to 12 class, and "
          "this is AI Trip Planner, a one-person AI travel agency.")
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
# decorative route: five stops in member colours joined by a dashed path
pts = [(8.9, 1.25), (10.25, 0.85), (11.55, 1.6), (10.6, 2.55), (12.05, 3.05)]
for (x1, y1), (x2, y2) in zip(pts, pts[1:]):
    c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb('3A4A6E'); c.line.width = Pt(2); c.line.dash_style = DASH.DASH
for (x, y), m in zip(pts, 'ABCDE'):
    badge(s, x - 0.22, y - 0.22, m, fill=MEMBERS[m][2], d=0.44, size=14, ring=True)
text(s, 0.9, 1.75, 8, 0.3, 'ELEC5620  ·  STAGE 1  ·  MON 10-12 GROUP 14', size=12, color='F2A65A', bold=True)
text(s, 0.9, 2.2, 9, 1.1, 'AI Trip Planner', size=60, color='FFFFFF', font=HEAD, bold=True)
box(s, 0.92, 3.42, 0.9, 0.06, fill=ORANGE, line=None, r=0)
text(s, 0.9, 3.7, 8.6, 0.8, 'A one-person AI travel agency: seven AI roles plan a trip together, one human runs the company.',
     size=18, color='C9D3E3')
cw = (MR - 0.9 - 4 * 0.18) / 5
for i, m in enumerate('ABCDE'):
    name, role, col = MEMBERS[m]
    x = 0.9 + i * (cw + 0.18)
    box(s, x, 5.05, cw, 1.45, fill='1C2B4D', line='2C3B60', r=0.12)
    box(s, x, 5.05, cw, 0.07, fill=col, line=None, r=0)
    badge(s, x + 0.2, 5.3, m, fill=col, d=0.36, size=12, ring=False)
    text(s, x + 0.66, 5.3, cw - 0.8, 0.36, name, size=14, bold=True, color='FFFFFF', anchor='m')
    text(s, x + 0.2, 5.8, cw - 0.4, 0.6, role, size=12, color='AEB9CC')

# ---------------------------------------------------------------- 2 problem
s = slide(2, 'A', 'The problem', 'Every travel choice constrains the others',
          "Planning a trip means juggling many sites, and every choice constrains the others. A chatbot "
          "can write an itinerary, but it invents prices and never checks that the parts fit. So we built an agency "
          "where specialist agents plan together on real data, and code checks their work.")
cards = [('layers', 'Too many sites', 'Flights, hotels, maps, food blogs and a budget spreadsheet, each on its own.', False),
         ('bot', 'Chatbots invent facts', 'One model writes a nice itinerary with made-up prices, and never checks that the parts fit.', False),
         ('sparkles', 'Our answer', 'Specialist agents plan together on real data, and code checks their work.', True)]
w = (MR - ML - 0.5) / 3
for i, (ic, h, b, hi) in enumerate(cards):
    x = ML + i * (w + 0.25)
    card(s, x, CT + 0.2, w, 4.6, fill=NAVY if hi else CARD)
    box(s, x + 0.4, CT + 0.6, 0.9, 0.9, fill=ORANGE if hi else TINT, line=None, r=0.45)
    icon(s, ic, x + 0.6, CT + 0.8, 0.5, 'FFFFFF' if hi else NAVY)
    text(s, x + 0.4, CT + 1.8, w - 0.8, 0.9, h, size=21, anchor='b', font=HEAD, bold=True, color='FFFFFF' if hi else NAVY)
    text(s, x + 0.4, CT + 2.85, w - 0.8, 1.3, b, size=15.5, color='C9D3E3' if hi else INK, spacing=1.25)
    if hi:
        text(s, x + 0.4, CT + 4.15, w - 0.8, 0.3, 'AI Trip Planner', size=12, bold=True, color='F2A65A')

# ---------------------------------------------------------------- 3 roles
s = slide(3, 'A', 'Users and AI roles', 'Two humans, seven AI roles',
          "There are two humans: the traveller, our customer, and the founder who runs the agency and tests "
          "the agents in our Agent Lab. Then seven AI roles: a coordinator that talks to the traveller, a supervisor "
          "that assigns work, and five specialists, each with its own prompt and tools.")
text(s, ML, CT, 3.0, 0.3, 'HUMANS', size=11, bold=True, color=MUTED)
for i, (ic, h, b) in enumerate((('user', 'Traveller', 'Describes the trip, reads and edits the plan'),
                                ('briefcase', 'Founder / operator', 'Runs the agency: models, providers, quotas, Agent Lab'))):
    y = CT + 0.4 + i * 2.35
    card(s, ML, y, 3.0, 2.15)
    box(s, ML + 0.3, y + 0.3, 0.62, 0.62, fill=TINT, line=None, r=0.31)
    icon(s, ic, ML + 0.43, y + 0.43, 0.36)
    text(s, ML + 0.3, y + 1.05, 2.6, 0.4, h, size=16, bold=True, color=NAVY, font=HEAD)
    text(s, ML + 0.3, y + 1.45, 2.5, 0.65, b, size=12, color=INK)
AX = ML + 3.3
text(s, AX, CT, 5, 0.3, 'AI ROLES', size=11, bold=True, color=MUTED)
aw = MR - AX
for i, (ic, h, b) in enumerate((('messages-square', 'Travel Coordinator', 'Chats with the traveller, builds the TripBrief, asks for missing facts'),
                                ('workflow', 'Planning Supervisor', 'Picks which specialists each step needs'))):
    y = CT + 0.4 + i * 1.12
    card(s, AX, y, aw, 0.92, fill=NAVY if i == 0 else CARD)
    icon(s, ic, AX + 0.3, y + 0.27, 0.38, 'FFFFFF' if i == 0 else NAVY)
    text(s, AX + 0.9, y + 0.12, 3.0, 0.68, h, size=16, bold=True, color='FFFFFF' if i == 0 else NAVY, anchor='m')
    text(s, AX + 3.8, y + 0.12, aw - 4.1, 0.68, b, size=13, color='C9D3E3' if i == 0 else INK, anchor='m')
specs = [('calendar-days', 'Itinerary', 'Day plan, route checks', 'B'), ('plane', 'Transport', 'Flights, rail, local routes', 'B'),
         ('bed-double', 'Accommodation', 'Stays, room allocation', 'C'), ('compass', 'Destination', 'Sights, customs, packing', 'D'),
         ('utensils', 'Dining', 'Venues, dietary needs', 'D')]
sw = (aw - 4 * 0.14) / 5
y = CT + 2.72
text(s, AX, y - 0.02, 5, 0.3, 'FIVE SPECIALISTS', size=11, bold=True, color=MUTED)
for i, (ic, h, b, m) in enumerate(specs):
    x = AX + i * (sw + 0.14)
    card(s, x, y + 0.36, sw, 2.12)
    box(s, x, y + 0.36, sw, 0.07, fill=MEMBERS[m][2], line=None, r=0)
    icon(s, ic, x + 0.2, y + 0.62, 0.42)
    text(s, x + 0.2, y + 1.18, sw - 0.28, 0.4, h, size=12.5, bold=True, color=NAVY)
    text(s, x + 0.2, y + 1.58, sw - 0.28, 0.8, b, size=11, color=INK)

# ---------------------------------------------------------------- 4 architecture
s = slide(4, 'A', 'Architecture', 'The graph drives the agents, not the other way round',
          "Our key design choice: the LangGraph workflow drives the agents, not the other way round. "
          "Specialists plan in stages on a shared board, code detects budget, time and route conflicts, and only the "
          "agents involved revise, for at most three rounds. If a model fails, a validated fallback keeps planning working.")
LW = 8.05
text(s, ML, CT, LW, 0.3, 'LANGGRAPH WORKFLOW', size=11, bold=True, color=MUTED)
figure(s, 'workflow.png', ML, CT + 0.35, LW, 1.75, callouts=[(1, 0.165, 0.12), (2, 0.335, 0.12), (3, 0.64, 0.02)])
notes4 = [('Plan in stages', 'Transport, then accommodation, then itinerary and dining, all on one shared board.'),
          ('Code finds conflicts', 'Budget, time and route conflicts are detected in code, not by a model.'),
          ('Targeted revision', 'Only the agents involved revise, at most three rounds, kept only if the plan improves.')]
for i, (h, b) in enumerate(notes4):
    note(s, i + 1, ML + 0.1, CT + 2.4 + i * 0.9, LW - 0.2, h, b, size=13)
text(s, ML + LW + 0.3, CT, 4, 0.3, 'REQUEST PATH', size=11, bold=True, color=MUTED)
figure(s, 'architecture.png', ML + LW + 0.3, CT + 0.35, MR - ML - LW - 0.3, CB - CT - 0.35)

# ---------------------------------------------------------------- 5 A template
behaviour_template(5, 'A')

# ---------------------------------------------------------------- 6 classification
s = slide(6, 'B', 'Requirement classification', 'Each member owns two core features and one optional',
          "We classified requirements into the agreed scope, mandatory "
          "capabilities and optional features. Each member owns two core features and one optional one, shown here. "
          "Every member also wrote an ad hoc requirement and broke it into functional, non-functional and constraint requirements.")
rows = [('A', 'Brief extraction and clarifying questions', 'Orchestration, conflicts, targeted revision', 'Agent Lab and Failure Lab'),
        ('B', 'Itinerary with route feasibility', 'Flights, rail and local transport', 'Traveller-chosen leg mode'),
        ('C', 'Accommodation, individual or group', 'Budget management', 'Keep a booked stay'),
        ('D', 'Destination guide', 'Dining with dietary needs', 'Weather-based packing'),
        ('E', 'Chat, timeline and map editing', 'Preferences and memory', 'Accounts and cloud sync')]
c0, c1, c2, c3 = ML, ML + 2.6, ML + 6.0, ML + 9.4
for x, lab in ((c1, 'CORE FEATURE 1'), (c2, 'CORE FEATURE 2'), (c3, 'OPTIONAL')):
    text(s, x + 0.15, CT, 3, 0.3, lab, size=11, bold=True, color=MUTED)
for i, (m, f1, f2, op) in enumerate(rows):
    y = CT + 0.38 + i * 0.7
    name, role, col = MEMBERS[m]
    card(s, ML, y, MR - ML, 0.6)
    badge(s, c0 + 0.18, y + 0.14, m, fill=col, ring=False, d=0.32, size=11)
    text(s, c0 + 0.62, y, 1.9, 0.6, name, size=13, bold=True, color=NAVY, anchor='m')
    text(s, c1 + 0.15, y, 3.2, 0.6, f1, size=13, anchor='m')
    text(s, c2 + 0.15, y, 3.2, 0.6, f2, size=13, anchor='m')
    pill(s, c3 + 0.15, y + 0.15, op, fill=TINT, fg=NAVY, size=11, bold=False, w=2.6)
y = CT + 4.05
card(s, ML, y, MR - ML, 1.1, fill=NAVY)
icon(s, 'scale', ML + 0.35, y + 0.32, 0.46, 'F2A65A')
text(s, ML + 1.1, y + 0.1, 3.2, 0.9, 'Models choose,\ncode decides.', size=20, font=HEAD, bold=True, color='FFFFFF', anchor='m')
text(s, ML + 4.4, y + 0.1, MR - ML - 4.7, 0.9,
     'Agreed scope: a single-user planner with AUD estimates; no booking or payment. No invented prices, no relaxed preferences.',
     size=14, color='C9D3E3', anchor='m')

# ---------------------------------------------------------------- 7 feature model
s = slide(7, 'B', 'Feature model', 'Seven feature groups with their constraints',
          "The feature diagram groups the system into seven areas, with mandatory, optional and alternative "
          "features, such as shared or individual rooms. We added cross-tree constraints and attached non-functional "
          "requirements, like money in integer cents.")
groups = [('messages-square', 'Conversation', ['● Chat intake', '● Brief extraction', '   ⊕ LLM | rule parser', '● Clarifying questions', '○ Attachments']),
          ('workflow', 'Multi-agent planning', ['● Supervisor delegation', '● Staged planning board', '● Conflict detection', '   ⊗ budget, time, geography, infeasible', '● Targeted revision ≤ 3']),
          ('boxes', 'Specialist roles', ['● Itinerary  ● Dining', '● Transport', '   ○ Traveller arranges flights', '● Accommodation', '   ⊕ shared | individual  ○ booked', '● Destination  ○ packing']),
          ('wallet', 'Budget', ['● Cost roll-up in AUD', '● Stage allocation', '● Minimum-cost floor']),
          ('shield-check', 'Evidence & tools', ['⊕ Live providers | mock fixtures', '● Provenance labels']),
          ('monitor', 'Workspace', ['⊕ Timeline | map', '● Plan editing', '● Preferences & filters', '○ Accounts & cloud sync', '○ Agent Lab']),
          ('cpu', 'Memory', ['● Session chat history', '● Long-term preferences', '⊕ Redis REST | in-process'])]
gw = (MR - ML - 3 * 0.18) / 4
gh = 2.05
for i, (ic, h, items) in enumerate(groups):
    x = ML + (i % 4) * (gw + 0.18)
    y = CT + (i // 4) * (gh + 0.18)
    card(s, x, y, gw, gh)
    icon(s, ic, x + 0.22, y + 0.2, 0.32, ORANGE)
    text(s, x + 0.66, y + 0.18, gw - 0.8, 0.36, h, size=14, bold=True, color=NAVY, anchor='m')
    text(s, x + 0.22, y + 0.68, gw - 0.34, gh - 0.75, items, size=11.5, color=INK, spacing=1.12, after=1)
x = ML + 3 * (gw + 0.18); y = CT + gh + 0.18
card(s, x, y, gw, gh, fill=NAVY)
text(s, x + 0.25, y + 0.2, gw - 0.45, gh - 0.3,
     [[('Notation', {'bold': True, 'color': 'F2A65A'})], '● mandatory   ○ optional', '⊕ exactly one   ⊗ one or more',
      [('Constraints', {'bold': True, 'color': 'F2A65A'})], 'Revision requires conflict detection; infeasible budget requires the floor; live providers require keys.'],
     size=10.5, color='DDE3EE', spacing=1.1, after=2)
text(s, ML, CB - 0.32, MR - ML, 0.3,
     [[('Non-functional:  ', {'bold': True, 'color': ORANGE}), ('schema checks at every boundary · money in integer cents · a provenance label on every section', {})]],
     size=12, color=INK)

# ---------------------------------------------------------------- 8 use cases
s = slide(8, 'B', 'Use cases', 'Ten use cases, one traveller, three external systems',
          "The use case diagram has one traveller and three external systems. Each member wrote at least "
          "one full specification, and our behaviour diagrams come from those.")
figure(s, 'use-case.png', ML, CT, 4.3, CB - CT)
ucs = [('A', 'Generate Itinerary', '«include» three use cases'),
       ('B', 'Arrange Transportation', 'Maps / Routes API'),
       ('C', 'Arrange Accommodation', 'Hotel search'),
       ('C', 'Manage Budget', '«extend» Edit Itinerary'),
       ('D', 'Weather clothing advice', 'Weather API'),
       ('D', 'Food recommendation', ''),
       ('E', 'Set Preferences · Submit Requirement · Edit · View', '')]
X = ML + 4.6
text(s, X, CT, 5, 0.3, 'WHO OWNS WHICH USE CASE', size=11, bold=True, color=MUTED)
for i, (m, uc, rel) in enumerate(ucs):
    y = CT + 0.4 + i * 0.68
    card(s, X, y, MR - X, 0.58)
    badge(s, X + 0.16, y + 0.13, m, fill=MEMBERS[m][2], ring=False, d=0.32, size=11)
    text(s, X + 0.62, y, 4.6, 0.58, uc, size=13.5, bold=True, color=NAVY, anchor='m')
    if rel:
        text(s, X + 5.0, y, MR - X - 5.15, 0.58, rel, size=11.5, color=MUTED, anchor='m', align='r')

# ---------------------------------------------------------------- 9 B template
behaviour_template(9, 'B')

# ---------------------------------------------------------------- 10 member C intro
s = slide(10, 'C', 'Accommodation and budget', 'From one ad hoc requirement to two use cases',
          "Member C owns accommodation and budget. The ad hoc requirement: the hotel must fit "
          "the money left after flights, a budget overrun should swap to a cheaper hotel that still meets the traveller's rules, and "
          "an impossible budget should name the minimum. That gave two use cases.")
card(s, ML, CT, MR - ML, 1.95, fill=NAVY)
icon(s, 'quote', ML + 0.35, CT + 0.3, 0.55, 'F2A65A')
text(s, ML + 1.2, CT + 0.25, MR - ML - 1.6, 1.3,
     'Flights get paid first, so the hotel has to fit in whatever money is left. If the trip ends up over budget, swap '
     'the hotel for a cheaper one that still meets my rules. If no hotel can ever fit, tell me the minimum I\'d need.',
     size=16.5, font=HEAD, italic=True, color='FFFFFF', spacing=1.2)
text(s, ML + 1.2, CT + 1.55, 6, 0.3, "AH-C1, in the traveller's words", size=11, bold=True, color='F2A65A')
uw = (MR - ML - 0.3) / 2
for i, (ic, h, items) in enumerate((('bed-double', 'UC-C1  Arrange Accommodation',
                                      ['Rooms from party size', 'Filter by rating and free cancellation', 'LLM picks one candidate id', 'Code validates and prices it in cents']),
                                     ('wallet', 'UC-C2  Manage Budget',
                                      ['Roll up costs in AUD', 'Spread an overrun over what each section can still cut', 'Revise only the sections asked to save', 'Report the minimum budget when nothing fits']))):
    x = ML + i * (uw + 0.3)
    y = CT + 2.25
    card(s, x, y, uw, CB - y, accent=ORANGE)
    box(s, x + 0.35, y + 0.3, 0.62, 0.62, fill=PEACH, line=None, r=0.31)
    icon(s, ic, x + 0.48, y + 0.43, 0.36, ORANGE)
    text(s, x + 1.15, y + 0.3, uw - 1.4, 0.62, h, size=17, bold=True, color=NAVY, font=HEAD, anchor='m')
    for j, it in enumerate(items):
        yy = y + 1.1 + j * 0.45
        icon(s, 'check', x + 0.4, yy + 0.04, 0.24, ORANGE)
        text(s, x + 0.8, yy, uw - 1.1, 0.35, it, size=14, color=INK, anchor='m')

# ---------------------------------------------------------------- 11 object diagram
s = slide(11, 'C', 'Complex structure', 'Object diagram: a Tokyo trip that is AUD 640 over budget',
          "The object diagram snapshots the planning board for a Tokyo trip that is 640 dollars over budget. "
          "The hotel prices are real fixture values; transport and dining are illustrative. Accommodation is asked to "
          "save 379 dollars, so it moves to the saver hotel.")
figure(s, 'object.png', ML, CT, MR - ML, 3.7, callouts=[(1, 0.29, 0.64), (2, 0.004, 0.55), (3, 0.885, 0.02)])
stats = [('1', 'Budget vs total', 'AUD 2,300 vs 2,940', 'over by AUD 640'),
         ('2', 'Accommodation asked to save', 'AUD 379.26', 'overrun spread by room to cut'),
         ('3', 'Hotel moves', 'standard → saver', 'AUD 1,520 → AUD 880')]
sw = (MR - ML - 0.4) / 3
for i, (n, h, big, sub) in enumerate(stats):
    x = ML + i * (sw + 0.2); y = CT + 3.9
    card(s, x, y, sw, CB - y)
    badge(s, x + 0.22, y + 0.2, n, ring=False, d=0.28, size=10.5)
    text(s, x + 0.62, y + 0.17, sw - 0.8, 0.34, h, size=12, bold=True, color=MUTED, anchor='m')
    text(s, x + 0.22, y + 0.55, sw - 0.4, 0.45, big, size=20, bold=True, color=NAVY, font=HEAD)
    text(s, x + 0.22, y + 0.97, sw - 0.4, 0.3, sub, size=11.5, color=MUTED)

# ---------------------------------------------------------------- 12 collaboration + structured class
s = slide(12, 'C', 'Complex structure', 'Collaboration and structured class',
          "The collaboration shows roles, not classes: the budget code plays the judge that turns proposals "
          "into targeted revisions. The structured class shows the workflow's parts and its injected ports.")
LW = 8.6
text(s, ML, CT, LW, 0.3, 'COLLABORATION  ·  NEGOTIATE TRIP PLAN', size=11, bold=True, color=MUTED)
figure(s, 'collaboration.png', ML, CT + 0.35, LW, 2.0, callouts=[(1, 0.355, 0.4)])
text(s, ML, CT + 2.6, LW, 0.3, 'STRUCTURED CLASS  ·  TRIPPLANNINGWORKFLOW', size=11, bold=True, color=MUTED)
figure(s, 'structured-class.png', ML, CT + 2.95, LW, CB - CT - 2.95, callouts=[(2, 0.47, 0.17), (3, 0.835, 0.02)])
X = ML + LW + 0.3
for i, (h, b) in enumerate((('Roles, not classes', 'My budget code plays the judge: it turns proposals into targeted revision requests.'),
                            ('Parts with multiplicity', 'The policy and budget calculator are parts of the workflow, one of each.'),
                            ('Injected ports', 'Plan, progress, tools, memory and model enter only through ports, so tests swap them.'))):
    note(s, i + 1, X, CT + 0.1 + i * 1.65, MR - X, h, b, size=13, h=1.5)

# ---------------------------------------------------------------- 13 activity
s = slide(13, 'C', 'Activity diagram · UC-C1', 'The LLM chooses once; code checks before any price',
          "In this activity diagram the LLM appears once, choosing one candidate id. Around it everything is "
          "deterministic: invalid output falls back to a rule, a booked stay skips the search, and prices are in cents.")
figure(s, 'c-activity.png', ML, CT, MR - ML, 4.05, callouts=[(1, 0.80, 0.70), (2, 0.31, 0.32), (3, 0.13, 0.08), (4, 0.67, 0.33)])
acts = [('LLM, once', 'Chooses one candidate id per stay.'),
        ('Guarded', 'Unknown id or off-schema output falls back to a deterministic pick.'),
        ('Booked stay', 'Skips the search entirely.'),
        ('Priced in cents', 'Plus the cheapest eligible total as the floor.')]
aw = (MR - ML - 0.6) / 4
for i, (h, b) in enumerate(acts):
    note(s, i + 1, ML + i * (aw + 0.2), CT + 4.3, aw, h, b, size=12.5, h=0.95)

# ---------------------------------------------------------------- 14 sequence
s = slide(14, 'C', 'Sequence diagram · UC-C2', 'Budget overrun and a targeted revision',
          "The sequence diagram shows an overrun: detectConflicts sends a revision only to the agents that "
          "can cut, and round two is kept only because the plan score improved.")
px, py, pw, ph = figure(s, 'c-sequence.png', ML, CT, 7.9, CB - CT, callouts=[(1, 0.88, 0.415), (2, 0.6, 0.53), (3, 0.62, 0.62), (4, 0.33, 0.77)])
X = ML + 8.2
for i, (h, b) in enumerate((('Round 1', 'The total is over budget.'),
                            ('detectConflicts', 'Sends a revision request with a target saving.'),
                            ('Round 2', 'Only the targeted agents revise.'),
                            ('Kept only if better', 'Otherwise the previous proposals stay. Infeasible: one conflict naming the minimum budget.'))):
    note(s, i + 1, X, CT + 0.1 + i * 1.25, MR - X, h, b, size=13, h=1.15)

# ---------------------------------------------------------------- 15 state machine
s = slide(15, 'C', 'State machine', 'The accommodation section within one planning turn',
          "The state machine follows the accommodation section through one turn, and maps its internal states "
          "onto the three statuses the traveller sees.")
figure(s, 'c-state.png', ML, CT, 5.4, CB - CT)
X = ML + 5.75
text(s, X, CT, 4, 0.3, 'INTERNAL STATE', size=11, bold=True, color=MUTED)
text(s, MR - 2.3, CT, 2.3, 0.3, 'TRAVELLER SEES', size=11, bold=True, color=MUTED, align='r')
maps = [(['Waiting', 'Planning', 'Kept', 'UnderReview', 'Revising'], 'planning', TINT, NAVY),
        (['Draft'], 'draft', 'E3F1E8', '2E7D4F'),
        (['NeedsYou'], 'needs_you', PEACH, ORANGE),
        (['Failed'], 'run stops', 'F1E4E4', '9B2C2C')]
y = CT + 0.4
for states, ui, bg, fg in maps:
    h = 1.55 if len(states) > 1 else 0.82
    card(s, X, y, MR - X, h)
    xx, yy = X + 0.25, y + 0.25
    for st in states:
        w = len(st) * 0.105 + 0.35
        if xx + w > MR - 2.6:
            xx, yy = X + 0.25, yy + 0.45
        pill(s, xx, yy, st, fill='FFFFFF', fg=NAVY, size=11.5, bold=False, line='B9C0CD', w=w, h=0.34)
        xx += w + 0.12
    icon(s, 'arrow-right', MR - 2.45, y + h / 2 - 0.13, 0.26, '9AA5B8')
    pill(s, MR - 1.95, y + h / 2 - 0.19, ui, fill=bg, fg=fg, size=12, w=1.7, h=0.38)
    y += h + 0.2

# ---------------------------------------------------------------- 16 class spine
s = slide(16, 'D', 'Elementary structure', 'Six class diagrams drawn from the code',
          "Our class model has six diagrams drawn from the code. The spine runs "
          "from the workspace, through the coordinator and workflow, to specialists that reach tools and memory only "
          "through injected interfaces.")
figure(s, 'class-spine.png', ML, CT, MR - ML, 4.15)
six = ['Structural spine', 'Domain contracts', 'Specialists', 'Ports and adapters', 'Use case trace', 'Generalisation']
w6 = (MR - ML - 5 * 0.15) / 6
for i, h in enumerate(six):
    x = ML + i * (w6 + 0.15); y = CT + 4.35
    on = i == 0
    card(s, x, y, w6, CB - y, fill=NAVY if on else CARD)
    text(s, x + 0.2, y + 0.12, 1, 0.4, str(i + 1), size=20, font=HEAD, bold=True, color='F2A65A' if on else MEMBERS['D'][2])
    text(s, x + 0.2, y + 0.5, w6 - 0.3, 0.35, h, size=12.5, bold=True, color='FFFFFF' if on else NAVY)

# ---------------------------------------------------------------- 17 specialists
s = slide(17, 'D', 'Elementary structure', 'One Specialist interface, five realisations',
          "Every agent realises one Specialist interface, so the graph loops over them without knowing their "
          "models. The registry aggregates agents, while a plan is composed of its sections.")
figure(s, 'class-specialists.png', ML, CT, MR - ML, 3.4, callouts=[(1, 0.395, 0.48), (2, 0.425, 0.355)])
rels = [('1', 'Specialist interface', 'realisation × 5', 'The graph loops over agents without knowing their models.'),
        ('2', 'Registry ◇ Specialist', 'aggregation 1 → 5', 'Agents live on their own; the registry only lists them.'),
        ('', 'TripPlan ◆ TripSection', 'composition 1 → 0..*', 'A section cannot outlive its plan.')]
rw = (MR - ML - 0.4) / 3
for i, (n, h, k, b) in enumerate(rels):
    x = ML + i * (rw + 0.2); y = CT + 3.6
    card(s, x, y, rw, CB - y)
    if n:
        badge(s, x + 0.22, y + 0.22, n, ring=False, d=0.28, size=10.5)
    text(s, x + (0.62 if n else 0.22), y + 0.18, rw - 0.8, 0.36, h, size=14, bold=True, color=NAVY, anchor='m')
    pill(s, x + 0.22, y + 0.62, k, fill=TINT, fg=NAVY, size=10.5, bold=False)
    text(s, x + 0.22, y + 1.05, rw - 0.4, 0.5, b, size=12, color=INK)

# ---------------------------------------------------------------- 18 rationale
s = slide(18, 'D', 'Design rationale', 'What we chose, and what we discarded',
          "We discarded a model-driven loop and free agent-to-agent chat, because limits and costs must be "
          "testable. We dropped approval checkpoints, since confirming applied nothing, and chose DeepSeek because it was "
          "about three times faster.")
dec = [('Who drives the loop', 'Deterministic LangGraph', 'Model-driven loop', 'limits and costs must be testable'),
       ('Collaboration', 'Staged board with budget shares', 'Free agent-to-agent chat', 'control flow would live in model output'),
       ('Human in the loop', 'Traveller edits; sections show needs_you', 'Approval checkpoints', 'confirming applied nothing'),
       ('LLM provider', 'DeepSeek', 'MiniMax', 'pages took 37–81 s against 15–23 s'),
       ('Prices', 'Tools only, with provenance', 'Model estimates', 'it invents hotels and fares'),
       ('Server state', 'Redis REST store', 'Process memory', 'lost on cold starts')]
c1, c2 = ML + 2.75, ML + 7.0
for x, lab, col in ((ML + 0.2, 'DECISION', MUTED), (c1 + 0.55, 'CHOSEN', '2E7D4F'), (c2 + 0.55, 'DISCARDED, AND WHY', '9B2C2C')):
    text(s, x, CT, 4, 0.3, lab, size=11, bold=True, color=col)
for i, (d, ch, dis, why) in enumerate(dec):
    y = CT + 0.38 + i * 0.79
    card(s, ML, y, MR - ML, 0.68)
    text(s, ML + 0.25, y, 2.4, 0.68, d, size=13, bold=True, color=NAVY, anchor='m')
    box(s, c1 + 0.08, y + 0.19, 0.3, 0.3, fill='E3F1E8', line=None, r=0.15)
    icon(s, 'check', c1 + 0.13, y + 0.24, 0.2, '2E7D4F')
    text(s, c1 + 0.55, y, 4.0, 0.68, ch, size=13, color=INK, anchor='m')
    box(s, c2 + 0.08, y + 0.19, 0.3, 0.3, fill='F1E4E4', line=None, r=0.15)
    icon(s, 'x', c2 + 0.13, y + 0.24, 0.2, '9B2C2C')
    text(s, c2 + 0.55, y, MR - c2 - 0.7, 0.68, [[(dis, {'bold': True, 'color': NAVY}), ('  ' + why, {'color': MUTED})]],
         size=12.5, anchor='m')

# ---------------------------------------------------------------- 19 D template
behaviour_template(19, 'D')

# ---------------------------------------------------------------- 20 package
s = slide(20, 'E', 'Package diagram', 'Dependencies point inward to the shared contracts',
          "Six packages, each with an owner, all depend inward on the shared "
          "contracts, which depend on nothing. New specialists, providers or stores plug in behind existing interfaces.")
figure(s, 'package.png', ML, CT, MR - ML, 3.3, callouts=[(1, 0.92, 0.47), (2, 0.025, 0.53)])
pts20 = [('package', 'Six packages', 'One monorepo, each package with an owner.', ''),
         ('target', 'shared: no dependencies', 'It holds the Zod contracts and ports.', '1'),
         ('monitor', 'Nothing imports web', 'So the planner runs and is tested without a browser.', '2'),
         ('boxes', 'Plug in at interfaces', 'A new specialist, provider or store.', '')]
pw4 = (MR - ML - 0.6) / 4
for i, (ic, h, b, n) in enumerate(pts20):
    x = ML + i * (pw4 + 0.2); y = CT + 3.5
    card(s, x, y, pw4, CB - y)
    icon(s, ic, x + 0.25, y + 0.2, 0.36, MEMBERS['E'][2])
    if n:
        badge(s, x + pw4 - 0.5, y + 0.22, n, ring=False, d=0.28, size=10.5)
    text(s, x + 0.25, y + 0.68, pw4 - 0.4, 0.35, h, size=13.5, bold=True, color=NAVY)
    text(s, x + 0.25, y + 1.06, pw4 - 0.4, 0.6, b, size=12, color=INK)

# ---------------------------------------------------------------- 21 deployment
s = slide(21, 'E', 'Deployment', 'One serverless app, optional cloud services',
          "It deploys as one serverless Next.js app. Memory, account sync and sign-in are optional, and "
          "without keys the same build runs fully on mock data.")
figure(s, 'deployment.png', ML, CT, 6.9, CB - CT, callouts=[(1, 0.03, 0.47), (2, 0.5, 0.35), (3, 0.83, 0.54)])
X = ML + 7.2
dep = [('1', 'Browser', 'Chats and trips in local storage; progress streams as NDJSON.'),
       ('2', 'Server', 'Next.js API routes and the orchestrator.'),
       ('3', 'Optional', 'Redis REST memory, Postgres account sync, Clerk sign-in.'),
       ('', 'No keys?', 'The same build runs offline on mock data.')]
for i, (n, h, b) in enumerate(dep):
    y = CT + i * 1.3
    if n:
        note(s, n, X, y + 0.1, MR - X, h, b, size=13.5, h=1.1)
    else:
        card(s, X, y, MR - X, 1.15, fill=NAVY)
        icon(s, 'wifi-off', X + 0.25, y + 0.38, 0.38, 'F2A65A')
        text(s, X + 0.85, y + 0.1, MR - X - 1.0, 0.95, [[(h + '  ', {'bold': True, 'color': 'F2A65A'}), (b, {})]],
             size=14, color='FFFFFF', anchor='m')

# ---------------------------------------------------------------- 22 change assessment
s = slide(22, 'E', 'Change assessment and acceptance', 'How a change gets in',
          "A change must trace to a requirement; a change to shared contracts needs a written decision. A pull "
          "request is accepted when CI passes, an end-to-end artifact exists, use case postconditions hold and the models are updated.")
cols = [('git-pull-request', 'Assess', ['Trace it to a requirement and use case', 'Shared contract touched? Write a decision note and tell every owner', 'Reversing a decision? Supersede it in writing']),
        ('clipboard-check', 'Accept', ['CI green: types, lint, tests, build, docs', 'End-to-end test leaves a repeatable artifact', 'Use case postconditions still hold', 'Docs and models updated in the same pull request'])]
cw = (MR - ML - 0.3) / 2
for i, (ic, h, items) in enumerate(cols):
    x = ML + i * (cw + 0.3); y = CT
    card(s, x, y, cw, 3.65, accent=MEMBERS['E'][2])
    box(s, x + 0.35, y + 0.3, 0.62, 0.62, fill=TINT, line=None, r=0.31)
    icon(s, ic, x + 0.48, y + 0.43, 0.36, MEMBERS['E'][2])
    text(s, x + 1.15, y + 0.3, 3, 0.62, h, size=20, font=HEAD, bold=True, color=NAVY, anchor='m')
    yy = y + 1.15
    for it in items:
        icon(s, 'check', x + 0.4, yy + 0.04, 0.24, MEMBERS['E'][2])
        text(s, x + 0.8, yy, cw - 1.1, 0.5, it, size=14, color=INK)
        yy += 0.8 if len(it) > 55 else 0.56
y = CT + 3.9
card(s, ML, y, MR - ML, CB - y, fill=NAVY)
text(s, ML + 0.35, y + 0.18, 4, 0.3, 'KEY ASSUMPTIONS', size=11, bold=True, color='F2A65A')
assume = ['One traveller per trip', 'AUD everywhere', 'Estimates, not bookings', 'Flights paid first', 'Models and providers can fail']
xx = ML + 0.35
for a in assume:
    w = pill(s, xx, y + 0.62, a, fill='243457', fg='FFFFFF', size=11.5, bold=False, h=0.4)
    xx += w + 0.12

# ---------------------------------------------------------------- 23 E template
behaviour_template(23, 'E')

# ---------------------------------------------------------------- 24 summary
s = slide(24, None, None, '', "In short: models choose, code decides, and the traveller stays in control. Thank you for watching.")
s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(NAVY)
text(s, 0.9, 1.3, 6, 0.3, 'SUMMARY', size=12, color='F2A65A', bold=True)
for i, (a, b) in enumerate((('Models', ' choose.'), ('Code', ' decides.'), ('The traveller', ' stays in control.'))):
    text(s, 0.9, 1.85 + i * 0.95, 11, 0.9, [[(a, {'color': 'F2A65A'}), (b, {})]], size=48, font=HEAD, bold=True, color='FFFFFF')
box(s, 0.92, 4.85, 0.9, 0.06, fill=ORANGE, line=None, r=0)
text(s, 0.9, 5.15, 11, 0.4, 'Stage 2: verify live providers end to end, extend the Agent Lab, and track the work in Jira.',
     size=17, color='C9D3E3')
text(s, 0.9, 5.75, 11, 0.4, 'github.com/Lilstanie/AI_TRIP_PLANNER', size=14, color='8A93A6')
for i, m in enumerate('ABCDE'):
    badge(s, MR - (5 - i) * 0.5, 6.55, m, fill=MEMBERS[m][2], d=0.38, size=12, ring=False)

out = sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, '..', 'ELEC5620_Stage1_Slides.pptx')
assert len(prs.slides) == TOTAL, len(prs.slides)
prs.save(out)
print('wrote', out)
